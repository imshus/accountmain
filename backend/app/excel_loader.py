"""Turn an uploaded Excel or CSV file into clean pandas DataFrames, one per sheet.

Accounting exports rarely start at cell A1: they carry a company name, report title and
blank rows above the real header, use "1,23,456.00", "(500)", "500-" or "1,200.00 Dr"
style numbers, DD/MM/YYYY dates, and a "Total" row at the bottom. This module finds the
header row, gives every column a SQL-friendly name, converts numbers and dates, and
drops total and balance rows so they are not counted twice.

Nothing is changed silently: every cell that could not be read, every duplicate row and
the comparison with the file's own Total row are reported as checks shown to the user.
"""

from __future__ import annotations

import csv
import functools
import io
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path

import numpy as np
import pandas as pd

from .compat import to_datetime
from .formatting import indian

# Reference-number columns: a Total row under them is usually a count of entries, not a sum.
_IDENTIFIER_NAME = re.compile(r"(^|_)(id|no|num|number|code|ref|sr|sl|serial|vou|voucher|vch|invoice|bill)(_|$)")


class UploadError(ValueError):
    """The file could not be used. The message is shown to the user."""


class _NoEntries(Exception):
    """A sheet has a header but no data rows. The message is shown to the user."""


@dataclass
class ParsedSheet:
    name: str
    frame: pd.DataFrame
    labels: dict[str, str]  # clean column name -> header text as written in the file
    header_row: int  # 1-based row number of the header within the sheet
    checks: list[dict] = field(default_factory=list)  # {"level": "ok" | "info" | "warning", "message": str}


@dataclass
class _Inferred:
    series: pd.Series
    kind: str  # "number" | "date" | "text"
    failed: list[int] = field(default_factory=list)  # positions of non-empty cells that could not be read
    notes: list[tuple[str, str]] = field(default_factory=list)  # (level, message)
    # (level, message with a "{rows}" placeholder, positions) for notes that name sheet rows
    row_notes: list[tuple[str, str, list[int]]] = field(default_factory=list)


# Share of non-empty cells that must parse for a column to count as numeric. Columns whose
# name says they hold money only need a majority, so a few bad cells are reported instead
# of turning the whole column into text.
_NUMERIC_SHARE = 0.8
_NUMERIC_SHARE_MONEY_NAME = 0.5
_DATE_SHARE = 0.8
_MONEY_NAME = re.compile(
    r"(^|_)(debit|credit|dr|cr|amount|amt|total|balance|bal|value|price|rate|qty|quantity|tax|gst|cgst|sgst|igst|"
    r"cess|tds|withdrawal|withdrawals|deposit|deposits|net|paid|received|receipt|payment|sales|purchase|cost|"
    r"expense|income|discount|freight)(_|$)"
)
# Debit/credit style columns: a "Dr"/"Cr" in their number format is a label, not a sign.
_DRCR_COLUMN = re.compile(r"(^|_)(debit|debits|credit|credits|dr|cr|withdrawal|withdrawals|deposit|deposits)(_|$)")
_MAX_ROWS_LISTED = 8

_CURRENCY = re.compile(r"(?i)(₹|rs\.?|inr|usd|eur|gbp|\$|€|£)")
_DR_CR = re.compile(r"(?i)^(?P<amount>.*\d\)?)\s*(?P<marker>dr|cr)\.?$")
_DR_CR_PREFIX = re.compile(r"(?i)^(?P<marker>dr|cr)\.?\s*(?P<amount>[\d(].*)$")
_DATE_LIKE = re.compile(
    r"^(\d{1,4}[-/.\s][A-Za-z0-9]{1,9}[-/.\s]\d{1,4}|[A-Za-z]{3,9}\.?\s+\d{1,2},?\s+\d{2,4})(\s+\d{1,2}:\d{2}(:\d{2})?)?$"
)
_DATE_NAME = re.compile(r"(^|_)(date|dt|dated)(_|$)")
_KEEP_AS_TEXT = re.compile(r"(^|_)(chq|cheque|utr|ref|reference|ifsc|acct|account_no|account_number|a_c_no|txn_id|transaction_id|upi_ref)(_|$)")
# Last word of a column name that marks it as holding identifiers (invoice, HSN, phone, account numbers...).
_ID_TOKENS = {
    "invoice", "inv", "bill", "vch", "voucher", "vou", "no", "nos", "num", "number", "code", "hsn", "sac", "pin",
    "pincode", "zip", "mobile", "mob", "phone", "ph", "tel", "telephone", "card", "acc", "acct", "gstin", "gstn",
    "pan", "utr", "cheque", "chq", "ref", "reference", "ifsc", "id",
}  # fmt: skip
_MEASURE_TOKENS = {"rate", "qty", "quantity", "amount", "amt", "value", "price", "per", "date", "dt", "dated", "total"}
_LEADING_ZERO = re.compile(r"^0\d+$")
_LONG_DIGITS = re.compile(r"^\+?\d{16,}$")
_MIXED_CODE = re.compile(r"^(?=[^\s]*\d)(?=[^\s]*[A-Za-z])[A-Za-z0-9]+(?:[-/][A-Za-z0-9]+)*$")
_WITH_UNIT = re.compile(
    r"(?i)^\d+(\.\d+)?\s*(g|gm|gms|gram|grams|kg|kgs|mg|ct|cts|carat|carats|pc|pcs|nos|no|ltr|lt|l|ml|m|mtr|cm|mm|ft|"
    r"dz|doz|box|bag|bags|k|lakh|lakhs|lac|cr|crore|crores)\.?$"
)

# Summary rows that must not be counted as entries.
_TOTAL_ROW = re.compile(
    r"(?i)^(?:(?:grand|sub|current|net|page|overall|final)[\s-]*)?totals?"
    r"(?:\s+(?:for|of)\b.*|\s*\(\s*[a-z0-9]{1,4}\s*\)|\s+[a-z0-9]{1,2}\b|\s*(?:rs\.?|inr|₹|amount|amt\.?))?\s*[:.\-=]*\s*$"
)
_CLOSING_ROW = re.compile(
    r"(?i)^(?:(?:to|by)\s+)?(?:closing\s+balance|balance\s+(?:c\s*/\s*[fd]|carried\s+(?:forward|down|over|fwd))|"
    r"c\s*/\s*[fd]|carried\s+(?:forward|over|fwd))\.?(?:\s*[:(].*)?\s*$"
)
_OPENING_ROW = re.compile(
    r"(?i)^(?:(?:to|by)\s+)?(?:opening\s+balance|balance\s+(?:b\s*/\s*[fd]|brought\s+(?:forward|down|over|fwd))|"
    r"b\s*/\s*[fd]|brought\s+(?:forward|over|fwd))\.?(?:\s*[:(].*)?\s*$"
)
_ROW_OPENING, _ROW_CLOSING, _ROW_TOTAL = 1, 2, 3
_SQL_RESERVED = {
    "all", "and", "as", "asc", "between", "by", "case", "check", "column", "create", "default",
    "delete", "desc", "distinct", "drop", "else", "end", "exists", "from", "group", "having", "in",
    "index", "insert", "into", "is", "join", "key", "like", "limit", "not", "null", "offset", "on",
    "or", "order", "primary", "references", "select", "set", "table", "then", "to", "union",
    "update", "values", "when", "where",
}  # fmt: skip


def read_workbook(content: bytes, filename: str) -> list[ParsedSheet]:
    ext = Path(filename).suffix.lower()
    workbook_checks: dict[str, list[dict]] = {}
    marks: dict[str, dict] = {}
    try:
        if ext == ".csv":
            raw_sheets = {Path(filename).stem: _read_csv(content)}
        elif ext == ".xls":
            raw_sheets, marks, workbook_checks = _read_xls(content, filename)
        else:
            raw_sheets, marks, workbook_checks = _read_xlsx(content)
    except UploadError:
        raise
    except Exception as exc:  # the readers raise many types for unreadable files
        raise UploadError(f"Could not read '{filename}': {exc}") from exc

    sheets, skipped = [], []
    for name, raw in raw_sheets.items():
        try:
            sheet = _parse_sheet(str(name), raw, marks.get(name))
        except _NoEntries as exc:
            skipped.append(str(exc))
            continue
        if sheet is not None:
            sheet.checks = workbook_checks.get(str(name), []) + sheet.checks
            sheets.append(sheet)
    if not sheets:
        raise UploadError(
            "No table found in the file. Each sheet needs a header row followed by at least one data row."
            + "".join(f" {message}" for message in skipped)
        )
    if skipped:
        sheets[0].checks = [{"level": "warning", "message": message} for message in skipped] + sheets[0].checks
    return sheets


def _split_format_sections(fmt: str) -> list[str]:
    """Split an Excel number format on ';' outside quotes and escapes (positive;negative;zero;text)."""
    out, current, quoted, i = [], [], False, 0
    while i < len(fmt):
        ch = fmt[i]
        if ch == "\\" and not quoted and i + 1 < len(fmt):
            current.append(fmt[i : i + 2])
            i += 2
            continue
        if ch == '"':
            quoted = not quoted
        if ch == ";" and not quoted:
            out.append("".join(current))
            current = []
        else:
            current.append(ch)
        i += 1
    out.append("".join(current))
    return out


_DRCR_IN_FORMAT = re.compile(r"(?i)(?<![a-z])(dr|cr)\.?(?![a-z])")
_FORMAT_LITERALS = re.compile(r'"[^"]*"|\\.|\[[^\]]*\]')


@functools.lru_cache(maxsize=4096)
def _format_sections(fmt: str) -> tuple[tuple[str | None, bool], ...]:
    """(Dr/Cr marker shown, whether it is a percent format) for each section of an Excel number format."""
    out = []
    for section in _split_format_sections(fmt):
        match = _DRCR_IN_FORMAT.search(section.replace("\\", ""))
        out.append((match.group(1).capitalize() if match else None, "%" in _FORMAT_LITERALS.sub("", section)))
    return tuple(out)


def _format_section(fmt: str, value: float) -> tuple[str | None, bool] | None:
    if not fmt or fmt == "General":
        return None
    sections = _format_sections(fmt)
    if value < 0:
        if len(sections) == 1:
            # One-section format: Excel shows the minus sign in front ("-250.00 Dr"), so the
            # marker does not decide the sign; leave the number as it is.
            return None, sections[0][1]
        return sections[1]
    return sections[0]


def _drcr_from_format(fmt: str | None, value: float) -> str | None:
    """'Dr' or 'Cr' when the cell's number format displays one (Tally exports keep the sign there)."""
    section = _format_section(fmt, value)
    return section[0] if section else None


def _number_cell(fmt: str | None, value):
    """Apply what an Excel number format says about a numeric cell.

    Returns (value, mark): a percent format (18% stored as 0.18) gives 18 and mark "%", so
    Excel and CSV files agree; a Dr/Cr format gives mark "Dr" or "Cr" with the value
    unchanged (whether the marker is a sign is decided per column, once the header is known).
    """
    section = _format_section(fmt, value)
    if section is None:
        return value, None
    marker, percent = section
    if percent:
        return round(value * 100, 10), "%"
    return value, marker


def _read_xlsx(content: bytes) -> tuple[dict[str, pd.DataFrame], dict[str, dict], dict[str, list[dict]]]:
    """Read every sheet cell by cell so that Dr/Cr shown only through number formats is kept.

    pandas.read_excel returns bare values, so a Tally balance displayed as "1,200.00 Cr"
    would be read as +1200. Here such cells are marked "Cr" and, if the column mixes Dr and
    Cr, become the text "1200 Cr", which the number parser turns into -1200. Also reports
    hidden sheets and formulas with no saved result.
    """
    from openpyxl import load_workbook

    values_wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    formulas_wb = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    raw_sheets, marks, checks = {}, {}, {}
    for ws in values_wb.worksheets:
        # The <dimension> tag written by some exporters is wrong; trusting it silently drops rows and columns.
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
        rows, sheet_marks = [], {}
        for r, row in enumerate(ws.iter_rows()):
            values = []
            for c, cell in enumerate(row):
                value = cell.value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    value, mark = _number_cell(getattr(cell, "number_format", None), value)
                    if mark:
                        sheet_marks[(r, c)] = mark
                values.append(value)
            rows.append(values)

        sheet_checks = []
        if ws.sheet_state != "visible":
            sheet_checks.append({"level": "info", "message": "This sheet is hidden in Excel; it was imported anyway."})
        missing = 0
        formulas_ws = formulas_wb[ws.title]
        if hasattr(formulas_ws, "reset_dimensions"):
            formulas_ws.reset_dimensions()
        for r, formula_row in enumerate(formulas_ws.iter_rows(values_only=True)):
            for c, formula in enumerate(formula_row):
                if isinstance(formula, str) and formula.startswith("=") and (
                    r >= len(rows) or c >= len(rows[r]) or rows[r][c] is None
                ):
                    missing += 1
        if missing:
            sheet_checks.append({
                "level": "warning",
                "message": f"{missing} formula cells have no saved result, so they were read as blank. Open the file "
                "in Excel, save it, and upload it again.",
            })

        width = max((len(r) for r in rows), default=0)
        raw_sheets[ws.title] = pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object)
        marks[ws.title] = sheet_marks
        checks[ws.title] = sheet_checks
    values_wb.close()
    formulas_wb.close()
    return raw_sheets, marks, checks


def _read_xls(content: bytes, filename: str) -> tuple[dict[str, pd.DataFrame], dict[str, dict], dict[str, list[dict]]]:
    """Read a .xls file, which in practice is one of several formats.

    Real Excel 97-2003 files are read with their number formats (Dr/Cr, percent), like .xlsx.
    Many bank and accounting downloads named .xls are really an HTML table or tab/comma
    separated text, and some are .xlsx files renamed; those are recognised by their content.
    """
    head = content[:8]
    if head.startswith(b"PK"):
        raw_sheets, marks, checks = _read_xlsx(content)
        for name in raw_sheets:
            checks.setdefault(name, []).insert(0, {
                "level": "info", "message": "This .xls file is really an .xlsx workbook; it was read as one.",
            })
        return raw_sheets, marks, checks
    if head.startswith(b"\xd0\xcf\x11\xe0") or head[:1] == b"\x09":
        return _read_biff(content)

    text = _decode_text(content)
    stripped = text.lstrip()
    lowered = stripped[:5000].lower()
    if stripped.startswith("<"):
        if "urn:schemas-microsoft-com:office:spreadsheet" in lowered and "<table" not in lowered:
            raise UploadError(
                f"'{filename}' is an 'XML Spreadsheet 2003' file. Open it in Excel, save it as .xlsx, and upload it again."
            )
        if "<table" in text.lower():
            tables = _read_html_tables(text)
            if not tables:
                raise UploadError(f"'{filename}' is a web page (HTML) with no table rows to read.")
            message = ("This .xls file is really a web page (HTML) table, as many bank and accounting downloads are; "
                       "it was read as a table.")
            return tables, {}, {name: [{"level": "info", "message": message}] for name in tables}
        raise UploadError(f"'{filename}' is not an Excel file (it holds web page or XML text without a table).")
    if "\x00" in stripped[:2000]:
        raise UploadError(f"Could not read '{filename}': it is not an Excel, HTML or text file.")
    name = Path(filename).stem
    message = ("This .xls file is really a text file with tab or comma separated columns, as many bank downloads are; "
               "it was read as one.")
    return {name: _read_csv(content)}, {}, {name: [{"level": "info", "message": message}]}


def _read_biff(content: bytes) -> tuple[dict[str, pd.DataFrame], dict[str, dict], dict[str, list[dict]]]:
    """Read an Excel 97-2003 workbook with xlrd, keeping number formats where the file has them."""
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content, formatting_info=True, on_demand=True)
        formats = True
    except (NotImplementedError, xlrd.XLRDError, AssertionError, KeyError, IndexError):
        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        formats = False
    raw_sheets, marks, checks = {}, {}, {}
    try:
        for index in range(book.nsheets):
            sheet = book.sheet_by_index(index)
            rows, sheet_marks = [], {}
            for r in range(sheet.nrows):
                values = []
                for c in range(sheet.ncols):
                    ctype, value = sheet.cell_type(r, c), sheet.cell_value(r, c)
                    if ctype == xlrd.XL_CELL_TEXT:
                        values.append(value)
                    elif ctype == xlrd.XL_CELL_NUMBER:
                        if float(value).is_integer():
                            value = int(value)
                        fmt = None
                        if formats:
                            try:
                                fmt = book.format_map[book.xf_list[sheet.cell_xf_index(r, c)].format_key].format_str
                            except (KeyError, IndexError, AttributeError):
                                fmt = None
                        value, mark = _number_cell(fmt, value)
                        if mark:
                            sheet_marks[(r, c)] = mark
                        values.append(value)
                    elif ctype == xlrd.XL_CELL_DATE:
                        try:
                            values.append(xlrd.xldate.xldate_as_datetime(value, book.datemode))
                        except (ValueError, OverflowError, xlrd.xldate.XLDateError):
                            values.append(value)
                    elif ctype == xlrd.XL_CELL_BOOLEAN:
                        values.append(bool(value))
                    else:  # empty, blank or an error value such as #N/A
                        values.append(None)
                rows.append(values)
            width = max((len(r) for r in rows), default=0)
            raw_sheets[sheet.name] = pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object)
            marks[sheet.name] = sheet_marks
            sheet_checks = []
            if getattr(sheet, "visibility", 0):
                sheet_checks.append({"level": "info", "message": "This sheet is hidden in Excel; it was imported anyway."})
            checks[sheet.name] = sheet_checks
            book.unload_sheet(index)
    finally:
        book.release_resources()
    return raw_sheets, marks, checks


class _TableParser(HTMLParser):
    """Collect the rows of every <table> in an HTML page (no third-party parser needed)."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str | None]]] = []
        self._stack: list[dict] = []  # open tables: {"rows": [...], "row": [...] | None, "cell": [...] | None, "span": int}

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            table = {"rows": [], "row": None, "cell": None, "span": 1}
            self._stack.append(table)
            self.tables.append(table["rows"])
        elif not self._stack:
            return
        elif tag == "tr":
            self._close_row()
            self._stack[-1]["row"] = []
        elif tag in ("td", "th"):
            top = self._stack[-1]
            self._close_cell()
            if top["row"] is None:
                top["row"] = []
            span = dict(attrs).get("colspan") or "1"
            top["cell"], top["span"] = [], int(span) if span.isdigit() and 0 < int(span) < 100 else 1
        elif tag == "br" and self._stack[-1]["cell"] is not None:
            self._stack[-1]["cell"].append(" ")

    def handle_endtag(self, tag):
        if not self._stack:
            return
        if tag in ("td", "th"):
            self._close_cell()
        elif tag == "tr":
            self._close_row()
        elif tag == "table":
            self._close_row()
            self._stack.pop()

    def handle_data(self, data):
        if self._stack and self._stack[-1]["cell"] is not None:
            self._stack[-1]["cell"].append(data)

    def _close_cell(self):
        top = self._stack[-1]
        if top["cell"] is not None:
            text = " ".join("".join(top["cell"]).split())
            top["row"].append(text or None)
            top["row"].extend([None] * (top["span"] - 1))
            top["cell"] = None

    def _close_row(self):
        top = self._stack[-1]
        self._close_cell()
        if top["row"] is not None:
            top["rows"].append(top["row"])
            top["row"] = None


def _read_html_tables(text: str) -> dict[str, pd.DataFrame]:
    parser = _TableParser()
    parser.feed(text)
    parser.close()
    out = {}
    for rows in parser.tables:
        filled = [r for r in rows if any(v is not None for v in r)]
        if len(filled) < 2:  # layout tables and one-line headings carry no entries
            continue
        width = max(len(r) for r in filled)
        out[f"Table {len(out) + 1}"] = pd.DataFrame([r + [None] * (width - len(r)) for r in filled], dtype=object)
    return out


def _decode_text(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16")
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1", errors="replace")


def _read_csv(content: bytes) -> pd.DataFrame:
    """Read a CSV whose rows may have different lengths (title lines above the header).

    pandas' reader fixes the column count from the first line, so a one-cell title row
    makes it reject the rest of the file; the csv module reads ragged rows fine.
    """
    text = _decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    width = max((len(row) for row in rows), default=0)
    return pd.DataFrame([row + [None] * (width - len(row)) for row in rows], dtype=object)


def _rows_text(rows: list[int]) -> str:
    shown = ", ".join(str(r) for r in rows[:_MAX_ROWS_LISTED])
    return f"{shown} and {len(rows) - _MAX_ROWS_LISTED} more" if len(rows) > _MAX_ROWS_LISTED else shown


def _fmt(value: float) -> str:
    return indian(value)


def _parse_sheet(name: str, raw: pd.DataFrame, marks: dict | None = None) -> ParsedSheet | None:
    raw = raw.map(_blank_to_nan).dropna(how="all").dropna(axis=1, how="all")
    if raw.empty:
        return None

    header_idx = _find_header_row(raw)
    header_row = int(header_idx) + 1
    body = raw.loc[raw.index > header_idx]
    if body.empty:
        raise _NoEntries(f'Sheet "{name}" has a header row (row {header_row}) but no entries under it, so it was not imported.')

    labels_in_file = [_header_text(value, i) for i, value in enumerate(raw.loc[header_idx])]
    names = _unique_names([clean_name(text) for text in labels_in_file])
    labels = dict(zip(names, labels_in_file))
    source_columns = list(raw.columns)
    body = body.set_axis(names, axis=1)

    checks: list[dict] = []
    if marks:
        body, format_checks = _apply_format_marks(body, marks, source_columns, names, labels, header_idx)
        checks += format_checks
    if header_row > 1:
        checks.append({"level": "info", "message": f"Header found on row {header_row}; rows above it were treated as titles."})

    # Each distinct cell read as a number once; reused by the type inference below.
    caches = {
        col: dict.fromkeys(pd.unique(body[col].dropna()), math.nan) if _identifier_name(col)
        else {v: to_number(v) for v in pd.unique(body[col].dropna())}
        for col in names
    }
    summary = _find_summary_rows(body, names, caches)
    total_rows = body.loc[summary.dropped]
    body = body.loc[~summary.dropped]
    checks += summary.checks
    if body.empty:
        raise _NoEntries(f'Sheet "{name}" has only total or balance rows under its header (row {header_row}), so it was not imported.')

    # Sheet row numbers (1-based) of the data rows, for pointing the user at problem cells.
    sheet_rows = [int(i) + 1 for i in body.index]
    section = None
    if summary.reference is not None and summary.section_start is not None:
        section = (body.index > summary.section_start) & (body.index < summary.reference)
    body = body.reset_index(drop=True)

    columns: dict[str, pd.Series] = {}
    for col in names:
        inferred = _infer_column(body[col], col, caches.get(col))
        columns[col] = inferred.series
        for level, note in inferred.notes:
            checks.append({"level": level, "message": f'"{labels[col]}": {note}'})
        for level, note, positions in inferred.row_notes:
            rows = [sheet_rows[p] for p in positions]
            checks.append({"level": level, "message": f'"{labels[col]}": ' + note.format(rows=_rows_text(rows))})
        if inferred.failed:
            rows = [sheet_rows[p] for p in inferred.failed]
            examples = ", ".join(f'"{body[col].iloc[p]}"' for p in inferred.failed[:3])
            checks.append({
                "level": "warning",
                "message": f'{len(rows)} cell{"s" if len(rows) > 1 else ""} in "{labels[col]}" could not be read as '
                f"{'numbers' if inferred.kind == 'number' else 'dates'} and {'were' if len(rows) > 1 else 'was'} left "
                f"blank (row {_rows_text(rows)}; e.g. {examples}). Fix them in the file and upload again.",
            })
    frame = pd.DataFrame(columns)

    reference = total_rows.loc[summary.reference] if summary.reference is not None else None
    checks += _reconcile_totals(frame, reference, labels, section)

    duplicated = frame.duplicated(keep="first").to_numpy()
    if duplicated.any():
        rows = [sheet_rows[p] for p in np.flatnonzero(duplicated)]
        checks.append({
            "level": "warning",
            "message": f"{len(rows)} row{'s are exact copies' if len(rows) > 1 else ' is an exact copy'} of an earlier row "
            f"(row {_rows_text(rows)}). {'They were' if len(rows) > 1 else 'It was'} kept; check it is not an entry made twice.",
        })

    return ParsedSheet(name=name, frame=frame, labels=labels, header_row=header_row, checks=checks)


def _apply_format_marks(body: pd.DataFrame, marks: dict, source_columns: list, names: list[str],
                        labels: dict[str, str], header_idx) -> tuple[pd.DataFrame, list[dict]]:
    """Turn Dr/Cr shown only by Excel number formats into signs, where the column mixes Dr and Cr.

    In a Debit or Credit column, or one where every amount shows the same marker, the marker
    labels the column rather than giving each amount a sign, so the amounts are left alone.
    """
    position = {col: i for i, col in enumerate(source_columns)}
    by_column: dict[str, list[tuple[int, str]]] = {}
    for (r, c), mark in marks.items():
        if r > header_idx and c in position and r in body.index:
            by_column.setdefault(names[position[c]], []).append((r, mark))

    checks, applied = [], 0
    for col in names:
        items = by_column.get(col)
        if not items:
            continue
        percent = sum(1 for _, mark in items if mark == "%")
        if percent:
            checks.append({
                "level": "info",
                "message": f'"{labels[col]}": {percent} cell{"s are" if percent > 1 else " is"} formatted as a percentage in '
                "Excel (0.18 shown as 18%) and were read as the percentage shown (18), as in a CSV file.",
            })
        drcr = [(r, mark) for r, mark in items if mark in ("Dr", "Cr")]
        if not drcr:
            continue
        sides = sorted({mark for _, mark in drcr})
        if _DRCR_COLUMN.search(col) or len(sides) == 1:
            reason = ("in a debit, credit, withdrawal or deposit column that is a label, not a sign"
                      if _DRCR_COLUMN.search(col) else f'every amount shows "{sides[0]}", which labels the column')
            checks.append({
                "level": "info",
                "message": f'"{labels[col]}": the Excel number format shows {" / ".join(sides)} after the amounts; '
                f"{reason}, so the amounts were kept as they are.",
            })
            continue
        index = [r for r, _ in drcr]
        body[col] = body[col].astype(object)
        current = body.loc[index, col]
        body.loc[index, col] = [f"{abs(value)} {mark}" if isinstance(value, (int, float)) else value
                                for value, (_, mark) in zip(current, drcr)]
        applied += len(index)
    if applied:
        checks.insert(0, {
            "level": "info",
            "message": f"{applied} amounts show Dr/Cr only through Excel number formatting (as Tally exports do); "
            "the Dr/Cr was read from the format.",
        })
    return body, checks


@dataclass
class _Summary:
    dropped: np.ndarray  # bool per body row
    checks: list[dict]
    reference: object = None  # index label of the total row to reconcile against
    section_start: object = None  # index label of the total row before it (or the header)


def _summary_kind(value) -> int:
    if not isinstance(value, str):
        return 0
    text = value.strip()
    if _TOTAL_ROW.match(text):
        return _ROW_TOTAL
    if _CLOSING_ROW.match(text):
        return _ROW_CLOSING
    if _OPENING_ROW.match(text):
        return _ROW_OPENING
    return 0


def _find_summary_rows(body: pd.DataFrame, names: list[str], caches: dict[str, dict]) -> _Summary:
    """Find total and balance rows, labelled or not, so they are not counted as entries.

    Labelled rows: "Total", "Grand Total", "Total for April", "Total (A)", "Closing Balance",
    "c/f", "b/f"... An "Opening Balance" (or "b/f") row is kept as an entry when it carries a date
    or comes before every other entry, as the opening line of a ledger does. Unlabelled rows:
    rows holding only numbers (no text, date or reference cell) that equal the sum of the
    entries since the previous total, of all entries so far, or the previous total plus the
    balance rows after it (Tally's final total).
    """
    n = len(body)
    dropped = np.zeros(n, dtype=bool)
    if n == 0:
        return _Summary(dropped, [])
    kind = np.zeros(n, dtype=np.int8)
    label_text: dict[int, str] = {}
    bad = np.zeros(n, dtype=bool)  # holds text, a date, or a reference number
    has_number = np.zeros(n, dtype=bool)
    has_date = np.zeros(n, dtype=bool)
    numeric = np.full((n, len(names)), np.nan)
    for j, col in enumerate(names):
        series = body[col]
        present = series.notna().to_numpy()
        if not present.any():
            continue
        cache = caches[col]
        values = series.map(cache).to_numpy(dtype="float64", na_value=np.nan) if not _identifier_name(col) else np.full(n, np.nan)
        numeric[:, j] = values
        kinds = {v: k for v in cache if (k := _summary_kind(v))}
        if kinds:
            col_kind = series.map(kinds).fillna(0).to_numpy(dtype=np.int8)
            for p in np.flatnonzero(col_kind > kind):
                label_text[p] = str(series.iloc[p]).strip()
            kind = np.maximum(kind, col_kind)
        dates = {v: True for v in cache if isinstance(v, (datetime, date))}
        if _DATE_NAME.search(col):
            has_date |= present
        elif dates:
            has_date |= series.map(dates).notna().to_numpy()
        is_number = present & ~np.isnan(values)
        has_number |= is_number
        bad |= present if (_identifier_name(col) or _DATE_NAME.search(col)) else (present & ~is_number)
    candidate = (kind == 0) & ~bad & has_number
    events = np.flatnonzero((kind > 0) | candidate)
    if events.size == 0:
        return _Summary(dropped, [])

    filled = np.nan_to_num(numeric)
    plain = np.ones(n, dtype=bool)
    plain[events] = False
    sums = np.vstack([np.zeros(len(names)), np.cumsum(filled * plain[:, None], axis=0)])
    counts = np.concatenate([[0], np.cumsum(plain)])

    section = np.zeros(len(names))
    grand = np.zeros(len(names))
    since_total = None  # previous total row plus every row after it, balance rows included
    section_rows = grand_rows = 0
    last = 0
    reference = section_start = None
    break_label = None  # index label of the last total row
    labels = body.index
    unlabelled, balances, kept_opening, totals = [], [], [], []
    for p in events:
        block, rows = sums[p] - sums[last], counts[p] - counts[last]
        section += block
        grand += block
        section_rows += rows
        grand_rows += rows
        if since_total is not None:
            since_total += block
        last = p + 1
        row = filled[p]
        is_total = kind[p] == _ROW_TOTAL
        matched_rule = None
        if kind[p] == 0:
            cells = ~np.isnan(numeric[p])
            values = numeric[p][cells]
            if np.any(values != 0):
                close = lambda target: np.allclose(values, target[cells], rtol=0, atol=0.005)  # noqa: E731
                if section_rows >= 2 and close(section):
                    matched_rule = "section"
                elif grand_rows >= 2 and close(grand):
                    matched_rule = "grand"
                elif since_total is not None and close(since_total):
                    matched_rule = "running"
            if matched_rule:
                unlabelled.append(p)
                is_total = True
        if is_total:
            dropped[p] = True
            if kind[p] == _ROW_TOTAL:
                totals.append(p)
            if matched_rule != "running":
                reference, section_start = labels[p], break_label
            break_label = labels[p]
            since_total = row.copy()
            section = np.zeros(len(names))
            section_rows = 0
            continue
        if kind[p] == _ROW_CLOSING or (kind[p] == _ROW_OPENING and not has_date[p] and grand_rows > 0):
            dropped[p] = True
            balances.append(p)
            if since_total is not None:
                since_total += row
            continue
        if kind[p] == _ROW_OPENING:
            kept_opening.append(p)
        # an ordinary entry (including an opening balance line)
        section += row
        grand += row
        section_rows += 1
        grand_rows += 1
        if since_total is not None:
            since_total += row

    checks = []
    sheet_row = lambda p: int(labels[p]) + 1  # noqa: E731
    skipped = np.flatnonzero(dropped)
    if skipped.size:
        count = int(skipped.size)
        if not unlabelled and not balances:
            message = (f"Skipped {count} total row{'s' if count > 1 else ''} (row {_rows_text([sheet_row(p) for p in skipped])}) "
                       "so amounts are not counted twice.")
        else:
            message = (f"Skipped {count} total and balance row{'s' if count > 1 else ''} "
                       f"(row {_rows_text([sheet_row(p) for p in skipped])}) so amounts are not counted twice.")
            if unlabelled:
                many = len(unlabelled) > 1
                message += (f" Row {_rows_text([sheet_row(p) for p in unlabelled])} {'have' if many else 'has'} no label but "
                            f"{'hold' if many else 'holds'} only the column totals of the rows above.")
            if balances:
                shown = ", ".join(f'{sheet_row(p)} ("{label_text.get(p, "")}")' for p in balances[:_MAX_ROWS_LISTED])
                many = len(balances) > 1
                message += f" Row {shown} {'are balance lines' if many else 'is a balance line'}, not {'entries' if many else 'an entry'}."
        checks.append({"level": "info", "message": message})
    for p in kept_opening:
        checks.append({
            "level": "info",
            "message": f'Row {sheet_row(p)} ("{label_text.get(p, "")}") was kept as an entry: it is the opening balance at the start of the entries.',
        })
    return _Summary(dropped, checks, reference, section_start)


def _reconcile_totals(frame: pd.DataFrame, grand: pd.Series | None, labels: dict[str, str], section=None) -> list[dict]:
    """Compare the file's own (last) Total row with what the rows add up to."""
    if grand is None:
        return []
    checks = []
    for col in frame.columns:
        if not pd.api.types.is_numeric_dtype(frame[col]) or pd.api.types.is_bool_dtype(frame[col]):
            continue
        stated = to_number(grand[col])
        if math.isnan(stated):
            continue
        values = frame[col].to_numpy(dtype="float64", na_value=np.nan)
        if _IDENTIFIER_NAME.search(col):
            entries = int((~np.isnan(values)).sum())
            if stated == entries:
                checks.append({"level": "ok", "message": f'The Total row counts {entries} entries under "{labels[col]}", matching the rows read.'})
            continue
        computed = math.fsum(values[~np.isnan(values)])
        difference = computed - stated
        if abs(difference) < 0.005:
            checks.append({
                "level": "ok",
                "message": f'"{labels[col]}" adds up to {_fmt(computed)}, matching the file\'s Total row.',
            })
            continue
        if section is not None and 0 < int(section.sum()) < len(frame):
            part = values[section]
            subtotal = math.fsum(part[~np.isnan(part)])
            if abs(subtotal - stated) < 0.005:
                checks.append({
                    "level": "ok",
                    "message": f'"{labels[col]}": the file\'s last total row ({_fmt(stated)}) is a subtotal of the rows '
                    "since the total before it, and matches them.",
                })
                continue
        checks.append({
            "level": "warning",
            "message": f'"{labels[col]}": the file\'s Total row says {_fmt(stated)} but the rows add up to '
            f"{_fmt(computed)} (difference {_fmt(difference)}). Some rows may be missing, duplicated or unreadable.",
        })
    return checks


def _blank_to_nan(value):
    if isinstance(value, str) and not value.strip():
        return np.nan
    return value


def _is_text_label(value) -> bool:
    return isinstance(value, str) and math.isnan(to_number(value))


def _is_header_label(value) -> bool:
    """Text, or a cell that is only a heading in a header: a real date (month columns) or a year (2024, 2025)."""
    if _is_text_label(value) or isinstance(value, (datetime, date)):
        return True
    if isinstance(value, (int, float, np.integer, np.floating)) and not isinstance(value, bool):
        return float(value).is_integer() and 1990 <= value <= 2100
    return False


def _has_data_cell(values) -> bool:
    return any(not _is_text_label(v) for v in values)


def _find_header_row(raw: pd.DataFrame):
    """First row (within the top 20) that is mostly filled and mostly labels.

    A row is passed over when the next row is also all text and at least as full while the
    data further down holds numbers or dates: it is then a title (e.g. "Cash Book | FY 2024-25")
    above the real header.
    """
    needed = max(1, math.ceil(raw.shape[1] * 0.5))
    top = [row.dropna() for _, row in raw.head(21).iterrows()]
    indexes = list(raw.index[: len(top)])
    for i, values in enumerate(top[:20]):
        if len(values) < needed:
            continue
        texts = sum(_is_text_label(v) for v in values)
        if not texts or sum(_is_header_label(v) for v in values) / len(values) < 0.6:
            continue
        if i + 1 < len(top):
            following = top[i + 1]
            if (
                len(following) >= len(values)
                and all(_is_text_label(v) for v in following)
                and any(_has_data_cell(rest) for rest in top[i + 2 :])
            ):
                continue
        return indexes[i]
    return raw.index[0]


def _header_text(value, position: int) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return f"Column {position + 1}"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip() or f"Column {position + 1}"


def clean_name(text: str) -> str:
    name = re.sub(r"[^0-9a-z]+", "_", text.lower()).strip("_") or "column"
    if name[0].isdigit():
        name = f"c_{name}"
    if name in _SQL_RESERVED:
        name = f"{name}_"
    return name[:60]


def _unique_names(names: list[str]) -> list[str]:
    seen: set[str] = set()
    out = []
    for name in names:
        candidate, n = name, 2
        while candidate in seen:
            candidate, n = f"{name}_{n}", n + 1
        seen.add(candidate)
        out.append(candidate)
    return out


# --- Type inference ----------------------------------------------------------------------


def to_number(value) -> float:
    """Read one cell as a number; NaN if it is not one.

    Handles 1,23,456.78 (Indian grouping), ₹/Rs./INR, (500) and 500- for negatives,
    "-" for zero, a trailing %, and Tally-style "1,200.00 Dr" / "1,200.00 Cr"
    (Dr is read as positive, Cr as negative).
    """
    if isinstance(value, bool):
        return math.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        number = float(value)
        return number if math.isfinite(number) else math.nan
    if not isinstance(value, str):
        return math.nan

    text = _CURRENCY.sub("", value).replace(",", "").replace(" ", "").replace("−", "-").strip()
    if text.endswith("/-"):  # "Rs 1000/-"
        text = text[:-2].strip()
    marker = None
    drcr = _DR_CR.match(text) or _DR_CR_PREFIX.match(text)
    if drcr:
        text = drcr.group("amount").strip()
        marker = drcr.group("marker").lower()
    text = text.replace(" ", "").rstrip("%")
    sign = 1.0
    if text.startswith("(") and text.endswith(")"):
        sign, text = -1.0, text[1:-1]
    elif len(text) > 1 and text.endswith("-") and text[:-1].replace(".", "", 1).isdigit():
        sign, text = -1.0, text[:-1]
    elif text.startswith("-") and len(text) > 1:
        sign, text = -1.0, text[1:]
    if marker:
        if sign < 0:
            return math.nan  # "(1,000.00) Cr" or "-500 Dr": two sign markers, ambiguous
        sign = -1.0 if marker == "cr" else 1.0
    if text in {"-", "–"}:  # accounting formats print zero as a dash
        return 0.0
    try:
        number = float(text)
    except ValueError:
        return math.nan
    if not math.isfinite(number):
        return math.nan
    return sign * number


# Dates typed by a user (add entry): one strict reading, never a guess.
_ENTRY_TIME = re.compile(r"^(?P<date>.+?)(?:[ T](?P<h>\d{1,2}):(?P<m>\d{2})(?::(?P<s>\d{2}))?)?$")
_ENTRY_YMD = re.compile(r"^(\d{4})([-/.])(\d{1,2})\2(\d{1,2})$")
_ENTRY_DMY = re.compile(r"^(\d{1,2})([-/.])(\d{1,2})\2(\d{4}|\d{2})$")
_ENTRY_NAMED = ["%d-%b-%Y", "%d %b %Y", "%d/%b/%Y", "%d-%B-%Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
                "%b %d %Y", "%B %d %Y", "%d-%b-%y", "%d %b %y"]
_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
           "November", "December"]
_DATE_HELP = "Enter it as DD/MM/YYYY (15/04/2025) or YYYY-MM-DD (2025-04-15)."


def parse_date_input(value) -> pd.Timestamp:
    """Read a date typed by a user strictly; raises ValueError with a message for the user.

    YYYY-MM-DD, YYYY/MM/DD and YYYY.MM.DD are year-month-day; DD/MM/YYYY, DD-MM-YYYY and
    DD.MM.YYYY (or a two-digit year) are day-month-year, the Indian order; "15-Apr-2025"
    style names are accepted. Month-first input such as 04/25/2025 and impossible dates
    are rejected instead of being read some other way.
    """
    if isinstance(value, (datetime, date)):
        return pd.Timestamp(value)
    text = str(value).strip()
    parts = _ENTRY_TIME.match(text)
    if not text or not parts:
        raise ValueError(f'"{text}" is not a date. {_DATE_HELP}')
    body = parts.group("date").strip()
    hour, minute, second = (int(parts.group(k) or 0) for k in ("h", "m", "s"))
    if hour > 23 or minute > 59 or second > 59:
        raise ValueError(f'"{text}" has an impossible time of day. {_DATE_HELP}')

    if m := _ENTRY_YMD.match(body):
        year, month, day = int(m.group(1)), int(m.group(3)), int(m.group(4))
        if month > 12:
            raise ValueError(f'"{text}" has no month {month}: a date starting with the year is read as year-month-day. {_DATE_HELP}')
    elif m := _ENTRY_DMY.match(body):
        day, month, year = int(m.group(1)), int(m.group(3)), int(m.group(4))
        if len(m.group(4)) == 2:
            year += 2000
        if month > 12:
            hint = " It looks like month/day/year (US order)." if day <= 12 else ""
            raise ValueError(f'"{text}" has no month {month}: dates are read day first.{hint} {_DATE_HELP}')
    else:
        for fmt in _ENTRY_NAMED:
            try:
                parsed = datetime.strptime(body.replace("Sept", "Sep"), fmt)
            except ValueError:
                continue
            return pd.Timestamp(parsed.replace(hour=hour, minute=minute, second=second))
        raise ValueError(f'"{text}" is not a date in a format that can be read without guessing. {_DATE_HELP}')
    if not 1900 <= year <= 2100:
        raise ValueError(f'"{text}" has the year {year}, which looks like a typing mistake. {_DATE_HELP}')
    try:
        return pd.Timestamp(datetime(year, month, day, hour, minute, second))
    except ValueError:
        raise ValueError(f'"{text}" is not a real date: {_MONTHS[month - 1]} {year} has no day {day}.') from None


def to_timestamp(value):
    """A date typed by a user as a Timestamp, or NaT when it is not a clear, valid date (see parse_date_input)."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return pd.NaT
    try:
        return parse_date_input(value)
    except ValueError:
        return pd.NaT


# One format is chosen per column (strict), day-first before month-first, as recommended for
# Indian data: pandas' dayfirst=True alone is not strict and can mix 04/05 and 05/04 readings.
_DATE_FORMATS = [
    "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y", "%d.%m.%y",
    "%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
    "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %b %y", "%d-%B-%Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S",
]
_MONTH_FIRST_FORMATS = ["%m/%d/%Y", "%m-%d-%Y", "%m/%d/%y"]
_AMBIGUOUS_DAY_MONTH = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-]\d{2,4}")


def _strptime(text: str, fmt: str):
    try:
        return pd.Timestamp(datetime.strptime(text.replace("Sept", "Sep"), fmt))
    except ValueError:
        return None


def _parse_dates(uniques) -> tuple[dict, list[tuple[str, str]]]:
    """Map each distinct cell to a Timestamp (or NaT) using the format that fits the column best."""
    texts = [v.strip() for v in uniques if isinstance(v, str) and v.strip()]
    best_fmt, best_hits = None, 0
    if texts:
        for fmt in _DATE_FORMATS + _MONTH_FIRST_FORMATS:
            hits = sum(_strptime(t, fmt) is not None for t in texts)
            if hits > best_hits:  # strictly greater: earlier (day-first) formats win ties
                best_fmt, best_hits = fmt, hits
    fallbacks = [f for f in _DATE_FORMATS if f != best_fmt]
    if best_fmt in _MONTH_FIRST_FORMATS:
        fallbacks = [f for f in _MONTH_FIRST_FORMATS if f != best_fmt] + fallbacks
    mapping = {}
    for value in uniques:
        if isinstance(value, (datetime, date)):
            mapping[value] = pd.Timestamp(value)
        elif isinstance(value, str) and value.strip():
            text = value.strip()
            parsed = _strptime(text, best_fmt) if best_fmt else None
            for fmt in fallbacks:
                if parsed is not None:
                    break
                parsed = _strptime(text, fmt)
            mapping[value] = parsed if parsed is not None else pd.NaT
        else:
            mapping[value] = pd.NaT

    notes = []
    if best_fmt in _MONTH_FIRST_FORMATS:
        notes.append(("warning", "dates were read as month/day/year (US order) because values such as "
                                 "04/25/2025 only fit that order. Check this is right."))
    elif best_fmt and best_fmt.startswith("%d") and best_fmt[2] in "/-." and best_fmt[3:5] == "%m":
        pairs = [_AMBIGUOUS_DAY_MONTH.match(t) for t in texts]
        if all(m and int(m.group(1)) <= 12 for m in pairs if m):
            notes.append(("info", "dates such as 04/05/2025 were read as day/month/year (4 May 2025), the Indian order."))
    return mapping, notes


def _swapped_day_month(uniques, stamps: dict) -> dict:
    """Real dates whose day and month Excel swapped, mapped to the corrected date.

    A DD/MM/YYYY file opened in Excel with US date settings turns 04/05/2025 (4 May) into the
    real date 5 April, while 13/05/2025 cannot be a US date and stays text. The pattern: every
    text date is DD/MM with a day above 12 and every real date has a day of 12 or less.
    """
    real = [v for v in uniques if isinstance(v, (datetime, date))]
    texts = [v.strip() for v in uniques if isinstance(v, str) and v.strip()]
    if not real or not texts:
        return {}
    for text in texts:
        match = _AMBIGUOUS_DAY_MONTH.match(text)
        if not match or int(match.group(1)) <= 12 or int(match.group(2)) > 12:
            return {}
    if any(pd.isna(stamps.get(v, pd.NaT)) for v in uniques if isinstance(v, str) and v.strip()):
        return {}
    if any(v.day > 12 for v in real):
        return {}
    return {v: pd.Timestamp(v).replace(month=v.day, day=v.month) for v in real}


def _text(value):
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _failed_positions(values: pd.Series, parsed: pd.Series, all_values: pd.Series) -> list[int]:
    """Positions (in the full column) of non-empty cells that did not parse."""
    failed_labels = values.index[parsed.isna().to_numpy()]
    return [all_values.index.get_loc(label) for label in failed_labels]


def _identifier_name(name: str) -> bool:
    """Whether a column's name says it holds identifiers (invoice, bill, voucher, HSN, PIN, phone, account no...)."""
    tokens = [t for t in name.split("_") if t and not t.isdigit()]
    if not tokens or _MEASURE_TOKENS.intersection(tokens):
        return False
    return tokens[-1] in _ID_TOKENS or bool(_KEEP_AS_TEXT.search(name))


def _code_value(uniques, numbers: dict):
    """A value showing that a mostly numeric column holds codes, with the reason; None if there is none."""
    for value in uniques:
        if isinstance(value, str):
            text = value.strip()
            if _LEADING_ZERO.match(text):
                return value, "has a leading zero"
            if _LONG_DIGITS.match(text):
                return value, "has more than 15 digits"
            if (_MIXED_CODE.match(text) and not _WITH_UNIT.match(text) and math.isnan(numbers.get(value, math.nan))
                    and not _DATE_LIKE.match(text)):
                return value, "mixes letters and digits"
        elif isinstance(value, (int, np.integer)) and not isinstance(value, bool) and abs(int(value)) >= 10**15:
            return value, "has more than 15 digits"
    return None


def _as_text(series: pd.Series) -> pd.Series:
    return series.map(lambda v: None if pd.isna(v) else _text(v)).astype("string")


def _infer_column(series: pd.Series, name: str, numbers: dict | None = None) -> _Inferred:
    values = series.dropna()
    money_name = bool(_MONEY_NAME.search(name))
    if values.empty:
        if money_name and not _identifier_name(name):
            # An amount column with no entries yet (e.g. Credit in a sheet of payments) stays a number column.
            return _Inferred(pd.Series(np.nan, index=series.index, dtype="float64"), "number",
                             notes=[("info", "is blank in every row; it was kept as an empty amount column.")])
        return _Inferred(pd.Series([None] * len(series), dtype="string"), "text")
    if _identifier_name(name):
        # Invoice, HSN, cheque, UTR, account and phone numbers are identifiers: as numbers they would lose
        # leading zeros (0011 -> 11) or digits beyond 15, and "105A" would be unreadable.
        return _Inferred(_as_text(series), "text")

    uniques = pd.unique(values)

    cache = numbers or {}
    numbers = {v: cache[v] if v in cache else to_number(v) for v in uniques}
    parsed = values.map(numbers)
    numeric_share = parsed.notna().mean()
    if numeric_share >= (_NUMERIC_SHARE_MONEY_NAME if money_name else _NUMERIC_SHARE):
        code = None if (money_name or _DATE_NAME.search(name)) else _code_value(uniques, numbers)
        if code is not None:
            example, reason = code
            position = next(i for i, v in enumerate(series) if v is example or (isinstance(v, type(example)) and v == example))
            return _Inferred(_as_text(series), "text", row_notes=[(
                "info",
                f'kept as text, not numbers, because it holds codes ("{_text(example)}" in row {{rows}} {reason}); '
                "every value is kept exactly as written.",
                [position],
            )])
        result = series.map(numbers).astype("float64")
        present = result.dropna().to_numpy()
        failed = _failed_positions(values, parsed, series)
        notes = []
        if any(isinstance(v, str) and _DR_CR.match(v.strip()) for v in uniques):
            notes.append(("info", "Dr/Cr amounts were read with Dr as positive and Cr as negative."))
        # Dates stored as Excel serial numbers (e.g. 45748 = 1 Apr 2025, 45748.5 = 1 Apr 2025 12:00) in a date-named column.
        if _DATE_NAME.search(name) and present.size and present.min() > 20000 and present.max() < 80000:
            dates = pd.to_datetime(result, unit="D", origin="1899-12-30")
            if np.all(np.mod(present, 1) == 0):
                return _Inferred(dates, "date", failed, [("info", "Excel date serial numbers were converted to dates.")])
            return _Inferred(dates.dt.round("s"), "date", failed, [(
                "info", "Excel date serial numbers with a time of day (e.g. 45748.5 = 1 Apr 2025 12:00) were converted "
                        "to dates and times.",
            )])
        if present.size and np.all(np.mod(present, 1) == 0) and np.abs(present).max() < 2**53:
            return _Inferred(result.astype("Int64"), "number", failed, notes)
        return _Inferred(result, "number", failed, notes)

    stamps, date_notes = _parse_dates(uniques)
    parsed = values.map(stamps)
    if parsed.notna().mean() >= _DATE_SHARE:
        failed = _failed_positions(values, parsed, series)
        row_notes = []
        swapped = _swapped_day_month(uniques, stamps)
        if swapped:
            stamps.update(swapped)
            positions = [i for i, v in enumerate(series) if isinstance(v, (datetime, date)) and v in swapped]
            before = series.iloc[positions[0]]
            row_notes.append((
                "warning",
                f"{len(positions)} date{'s' if len(positions) > 1 else ''} had day and month swapped, as happens when a "
                "DD/MM/YYYY file is opened in Excel with US date settings (every date typed as text has a day above 12, "
                "every real date a day of 12 or less). Their day and month were swapped back (row {rows}; e.g. "
                f"{pd.Timestamp(before):%d %b %Y} became {swapped[before]:%d %b %Y}). Check these dates.",
                positions,
            ))
        return _Inferred(to_datetime(series.map(stamps)), "date", failed, date_notes, row_notes)

    text = _as_text(series)
    if money_name and numeric_share > 0:
        return _Inferred(text, "text", notes=[(
            "warning",
            f"looks like an amount column, but only {numeric_share:.0%} of its cells are numbers, so it was kept as "
            "text and is not included in totals.",
        )])
    return _Inferred(text, "text")
